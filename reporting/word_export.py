"""
Markdown report -> .docx, because the brief asks for MS Word output.

Deliberately a small converter for the subset of Markdown our reports actually
use (headings, bullets, tables, bold, links) rather than a general-purpose one —
the input is always our own generated report, so the shape is known.
"""
import io
import re

HEADING_RE = re.compile(r"^(#{1,6})\s+(.*)$")
BULLET_RE = re.compile(r"^[-*]\s+(.*)$")
LINK_RE = re.compile(r"\[([^\]]*)\]\(([^)]+)\)")
BOLD_RE = re.compile(r"\*\*(.+?)\*\*")
ITALIC_RE = re.compile(r"(?<!\*)\*(?!\*)(.+?)(?<!\*)\*(?!\*)")


def _plain(text: str) -> str:
    """Flatten inline markdown. Links become 'label (url)' so the URL survives
    in a printed document, where a hyperlink would otherwise be invisible."""
    text = LINK_RE.sub(lambda m: m.group(1) if m.group(1) == m.group(2) else f"{m.group(1)} ({m.group(2)})", text)
    text = BOLD_RE.sub(r"\1", text)
    text = ITALIC_RE.sub(r"\1", text)
    return text.strip()


def _is_table_row(line: str) -> bool:
    return line.strip().startswith("|") and line.strip().endswith("|")


def _is_separator_row(line: str) -> bool:
    return bool(re.fullmatch(r"\|[\s:|-]+\|", line.strip()))


def _split_row(line: str) -> list[str]:
    return [c.strip() for c in line.strip().strip("|").split("|")]


def markdown_to_docx_bytes(markdown: str, title: str = "Report") -> bytes:
    from docx import Document
    from docx.shared import Pt

    doc = Document()
    doc.core_properties.title = title

    lines = markdown.split("\n")
    i = 0
    while i < len(lines):
        line = lines[i]
        stripped = line.strip()

        if not stripped:
            i += 1
            continue

        # Table: consume the whole block at once
        if _is_table_row(stripped):
            block = []
            while i < len(lines) and _is_table_row(lines[i].strip()):
                if not _is_separator_row(lines[i]):
                    block.append(_split_row(lines[i]))
                i += 1
            if block:
                cols = max(len(r) for r in block)
                table = doc.add_table(rows=0, cols=cols)
                table.style = "Light Grid Accent 1"
                for r_idx, row in enumerate(block):
                    cells = table.add_row().cells
                    for c_idx in range(cols):
                        text = _plain(row[c_idx]) if c_idx < len(row) else ""
                        cells[c_idx].text = text
                        if r_idx == 0:
                            for p in cells[c_idx].paragraphs:
                                for run in p.runs:
                                    run.bold = True
                doc.add_paragraph()
            continue

        heading = HEADING_RE.match(stripped)
        if heading:
            level = len(heading.group(1))
            doc.add_heading(_plain(heading.group(2)), level=min(level, 4))
            i += 1
            continue

        bullet = BULLET_RE.match(stripped)
        if bullet:
            doc.add_paragraph(_plain(bullet.group(1)), style="List Bullet")
            i += 1
            continue

        para = doc.add_paragraph(_plain(stripped))
        if stripped.startswith("*") and stripped.endswith("*"):
            for run in para.runs:
                run.italic = True
                run.font.size = Pt(9)
        i += 1

    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()
